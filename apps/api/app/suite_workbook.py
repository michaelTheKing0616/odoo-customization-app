"""Configuration Suite: setup workbook (stage 1).

Download a workbook prefilled with what the database already has; fill it in;
upload it back. Code diffs each sheet against the live database, so an
unchanged workbook produces zero steps. Values are read as typed; nothing is
guessed.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook, load_workbook

META = "_ingenium"
YES, NO = "Yes", "No"
MAX_ROWS = 2000


def _x(client: Any, model: str, method: str, args: list[Any] | None = None, kw: dict[str, Any] | None = None) -> Any:
    return client.execute_kw(model, method, args or [], kw or {})


def _try(fn: Any, default: Any) -> Any:
    try:
        return fn()
    except Exception:  # noqa: BLE001  app not installed or no access
        return default


def _m2o(v: Any) -> str:
    return str(v[1]) if isinstance(v, list | tuple) and len(v) == 2 else ""


def _num(v: Any) -> float:
    try:
        return round(float(v), 6)
    except (TypeError, ValueError):
        return 0.0


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def setting_fields(client: Any) -> list[dict[str, Any]]:
    """Plain on/off and choice settings this database has (no installs, no accounting)."""
    from app.suite_setup import _GUARDED_SETTING

    meta = _try(lambda: _x(client, "res.config.settings", "fields_get", [], {"attributes": ["type", "string", "selection", "readonly"]}), {})
    out = []
    for f, m in sorted(meta.items()):
        if m.get("readonly") or f.startswith("module_") or (_GUARDED_SETTING.search(f) and not f.startswith("group_")):
            continue
        if m.get("type") in {"boolean", "selection"}:
            out.append({"field": f, "label": m.get("string") or f, "type": m["type"], "selection": m.get("selection") or []})
    return out


def live_snapshot(client: Any) -> dict[str, list[dict[str, Any]]]:
    sf = setting_fields(client)
    cur = _try(lambda: _x(client, "res.config.settings", "default_get", [[s["field"] for s in sf]]), {}) if sf else {}
    settings = [{"field": s["field"], "label": s["label"],
                 "value": (YES if cur.get(s["field"]) else NO) if s["type"] == "boolean" else _cell(cur.get(s["field"]) or "")}
                for s in sf]
    umeta = _try(lambda: _x(client, "uom.uom", "fields_get", [], {"attributes": ["type"]}), {})
    units: list[dict[str, Any]] = []
    if "relative_uom_id" in umeta:
        for u in _try(lambda: _x(client, "uom.uom", "search_read", [[]], {"fields": ["name", "relative_uom_id", "relative_factor"], "limit": MAX_ROWS}), []):
            units.append({"name": u["name"], "reference": _m2o(u.get("relative_uom_id")) or u["name"], "factor": _num(u.get("relative_factor") or 1)})
    elif umeta:
        rows = _try(lambda: _x(client, "uom.uom", "search_read", [[]], {"fields": ["name", "category_id", "factor", "uom_type"], "limit": MAX_ROWS}), [])
        refs = {_m2o(u["category_id"]): u["name"] for u in rows if u.get("uom_type") == "reference"}
        for u in rows:
            f = _num(u.get("factor") or 1)
            units.append({"name": u["name"], "reference": refs.get(_m2o(u["category_id"]), u["name"]), "factor": _num(1.0 / f) if f else 0.0})
    cats = [{"name": c["name"], "parent": _m2o(c.get("parent_id"))}
            for c in _try(lambda: _x(client, "product.category", "search_read", [[]], {"fields": ["name", "parent_id"], "limit": MAX_ROWS}), [])]
    prods = [{"name": p["name"], "default_code": p.get("default_code") or "", "category": _m2o(p.get("categ_id")),
              "list_price": _num(p.get("list_price")), "unit": _m2o(p.get("uom_id"))}
             for p in _try(lambda: _x(client, "product.template", "search_read", [[]],
                                      {"fields": ["name", "default_code", "categ_id", "list_price", "uom_id"], "limit": MAX_ROWS, "order": "id"}), [])]
    curr = [{"code": c["name"], "active": YES if c.get("active") else NO}
            for c in _try(lambda: _x(client, "res.currency", "search_read", [[]],
                                     {"fields": ["name", "active"], "context": {"active_test": False}, "order": "name"}), [])]
    return {"settings": settings, "units": units, "categories": cats, "products": prods, "currencies": curr}


SHEETS: list[tuple[str, str, list[tuple[str, str]]]] = [
    ("settings", "Settings", [("field", "Setting"), ("label", "What it is"), ("value", "Value")]),
    ("units", "Units", [("name", "Unit"), ("reference", "Measured against"), ("factor", "How many of that unit")]),
    ("categories", "Product categories", [("name", "Category"), ("parent", "Inside category")]),
    ("products", "Products", [("name", "Product"), ("default_code", "Internal reference"), ("category", "Category"),
                              ("list_price", "Sales price"), ("unit", "Unit")]),
    ("currencies", "Currencies", [("code", "Currency"), ("active", "In use")]),
]


def build_workbook(client: Any, *, connection_id: str, server_version: str) -> bytes:
    snap = live_snapshot(client)
    wb = Workbook()
    wb.remove(wb.active)
    for key, title, cols in SHEETS:
        ws = wb.create_sheet(title)
        ws.append([c[1] for c in cols])
        for row in snap[key]:
            ws.append([row.get(c[0], "") for c in cols])
        ws.freeze_panes = "A2"
    meta = wb.create_sheet(META)
    meta.append(["connection_id", connection_id])
    meta.append(["odoo_version", server_version])
    meta.append(["made_at", datetime.now(UTC).isoformat(timespec="seconds")])
    meta.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_workbook(data: bytes) -> tuple[dict[str, list[dict[str, str]]], dict[str, str]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: dict[str, list[dict[str, str]]] = {}
    for key, title, cols in SHEETS:
        if title not in wb.sheetnames:
            continue
        rows = list(wb[title].iter_rows(values_only=True))
        out[key] = []
        for r in rows[1:MAX_ROWS + 1]:
            rec = {c[0]: _cell(r[i] if i < len(r) else "") for i, c in enumerate(cols)}
            if any(rec.values()):
                out[key].append(rec)
    meta: dict[str, str] = {}
    if META in wb.sheetnames:
        for r in wb[META].iter_rows(values_only=True):
            if r and r[0]:
                meta[str(r[0])] = _cell(r[1] if len(r) > 1 else "")
    return out, meta


def _f(v: str) -> float | None:
    try:
        return round(float(v), 6)
    except ValueError:
        return None


def workbook_steps(client: Any, data: bytes) -> dict[str, Any]:
    """Diff the uploaded workbook against the live database. Unchanged rows give no steps."""
    sheets, meta = read_workbook(data)
    snap = live_snapshot(client)
    steps: list[dict[str, Any]] = []
    notes: list[str] = []
    src = {"kind": "workbook"}

    def add(step: dict[str, Any], sheet: str, row: int) -> None:
        step["id"] = f"w{len(steps) + 1}"
        step["source"] = {**src, "ref": f"{sheet} row {row}"}
        steps.append(step)

    sf = {s["field"]: s for s in setting_fields(client)}
    live_set = {s["field"]: s["value"] for s in snap["settings"]}
    for i, r in enumerate(sheets.get("settings", []), start=2):
        f, v = r["field"], r["value"]
        if f not in sf:
            notes.append(f"Settings row {i}: '{f}' is not a setting this suite changes here. Skipped.")
            continue
        if v == live_set.get(f):
            continue
        if sf[f]["type"] == "boolean":
            if v not in {YES, NO}:
                notes.append(f"Settings row {i}: write Yes or No.")
                continue
            val: Any = v == YES
        else:
            keys = [k for k, _ in sf[f]["selection"]]
            if v not in keys:
                notes.append(f"Settings row {i}: choose one of {', '.join(keys)}.")
                continue
            val = v
        add({"type": "setting", "inputs": {"values": {f: val}}}, "Settings", i)

    live_units = {u["name"]: u for u in snap["units"]}
    for i, r in enumerate(sheets.get("units", []), start=2):
        if r["name"] in live_units:
            lu = live_units[r["name"]]
            if r["reference"] != lu["reference"] or _f(r["factor"]) != _f(str(lu["factor"])):
                notes.append(f"Units row {i}: '{r['name']}' exists. Changing a unit in use is done on Odoo's Units screen. Skipped.")
            continue
        fac = _f(r["factor"])
        if not r["reference"] or fac is None or fac <= 0:
            notes.append(f"Units row {i}: give the unit it is measured against and a number above zero.")
            continue
        add({"type": "units", "inputs": {"name": r["name"], "reference": r["reference"], "factor": fac}}, "Units", i)

    live_cats = {c["name"] for c in snap["categories"]}
    seen_c: dict[str, int] = {}
    for i, r in enumerate(sheets.get("categories", []), start=2):
        if r["name"] in seen_c:
            notes.append(f"Product categories row {i}: '{r['name']}' is already on row {seen_c[r['name']]}. Skipped.")
            continue
        seen_c[r["name"]] = i
        if r["name"] in live_cats:
            continue
        vals: dict[str, Any] = {"name": r["name"]}
        if r["parent"]:
            vals["parent_id"] = {"ref": r["parent"], "model": "product.category"}
        add({"type": "create", "inputs": {"model": "product.category", "values": vals}}, "Product categories", i)

    live_p = {p["name"]: p for p in snap["products"]}
    seen_p: dict[str, int] = {}
    for i, r in enumerate(sheets.get("products", []), start=2):
        if r["name"] in seen_p:
            notes.append(f"Products row {i}: '{r['name']}' is already on row {seen_p[r['name']]}. Keep one row per product. Skipped.")
            continue
        seen_p[r["name"]] = i
        price = _f(r["list_price"]) if r["list_price"] else None
        lp = live_p.get(r["name"])
        if lp:
            vals = {}
            if price is not None and price != lp["list_price"]:
                vals["list_price"] = price
            if r["default_code"] != lp["default_code"]:
                vals["default_code"] = r["default_code"] or False
            if r["category"] and r["category"] != lp["category"]:
                vals["categ_id"] = {"ref": r["category"], "model": "product.category"}
            if vals:
                add({"type": "update", "inputs": {"model": "product.template", "match": {"name": r["name"]}, "values": vals}}, "Products", i)
            continue
        vals = {"name": r["name"]}
        if r["default_code"]:
            vals["default_code"] = r["default_code"]
        if price is not None:
            vals["list_price"] = price
        if r["category"]:
            vals["categ_id"] = {"ref": r["category"], "model": "product.category"}
        if r["unit"]:
            vals["uom_id"] = {"ref": r["unit"], "model": "uom.uom"}
        add({"type": "create", "inputs": {"model": "product.template", "values": vals}}, "Products", i)

    live_c = {c["code"]: c["active"] for c in snap["currencies"]}
    for i, r in enumerate(sheets.get("currencies", []), start=2):
        if r["code"] not in live_c:
            notes.append(f"Currencies row {i}: Odoo has no currency {r['code']}.")
            continue
        if r["active"] != live_c[r["code"]] and r["active"] in {YES, NO}:
            add({"type": "update", "task": "2.5#61", "inputs": {"model": "res.currency", "match": {"name": r["code"], "active": r["active"] != YES},
                                                                 "values": {"active": r["active"] == YES}}}, "Currencies", i)
    return {"steps": steps, "notes": notes, "meta": meta}
